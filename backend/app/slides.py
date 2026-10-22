"""Extract readable text from a PowerPoint .pptx file, one block per slide."""

import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class SlideError(ValueError):
    """The file is not a readable PowerPoint presentation."""


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text).strip()


def _shape_lines(shape) -> list[str]:
    lines: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            lines.extend(_shape_lines(child))
        return lines
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
        return lines
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            lines.append(text)
    return lines


def extract_slides(path: Path) -> tuple[int, list[tuple[int, str]]]:
    """Return (slide count, [(1-based slide number, text)] for slides that have text)."""
    try:
        deck = Presentation(str(path))
    except Exception as e:
        raise SlideError("The file is not a valid PowerPoint (.pptx).") from e
    slides: list[tuple[int, str]] = []
    try:
        for number, slide in enumerate(deck.slides, start=1):
            parts: list[str] = []
            for shape in slide.shapes:
                parts.extend(_shape_lines(shape))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    parts.append(f"Speaker notes: {notes}")
            text = _clean("\n".join(parts))
            if text:
                slides.append((number, text))
        return len(deck.slides), slides
    except SlideError:
        raise
    except Exception as e:
        raise SlideError("The presentation could not be read; it may be corrupted.") from e
